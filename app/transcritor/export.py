"""Transcript exports: plain text, SRT subtitles and Word documents."""

import io

from docx import Document
from docx.shared import Pt


def speaker_name(meta: dict, speaker: str) -> str:
    return meta.get("speaker_names", {}).get(speaker) or f"Falante {speaker}"


def _ts(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _srt_ts(seconds: float) -> str:
    ms = round(seconds * 1000)
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d},{ms % 1000:03d}"


def _sentence_text(sentence: dict) -> str:
    return " ".join(w["text"] for w in sentence["words"])


def _turn_text(turn: dict) -> str:
    return " ".join(_sentence_text(s) for s in turn["sentences"])


def to_txt(meta: dict, data: dict) -> bytes:
    lines = [meta["title"], ""]
    for turn in data["turns"]:
        lines.append(f"[{_ts(turn['start'])}] {speaker_name(meta, turn['speaker'])}: {_turn_text(turn)}")
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def to_srt(meta: dict, data: dict) -> bytes:
    blocks = []
    n = 0
    for turn in data["turns"]:
        name = speaker_name(meta, turn["speaker"])
        for s in turn["sentences"]:
            n += 1
            blocks.append(f"{n}\n{_srt_ts(s['start'])} --> {_srt_ts(s['end'])}\n{name}: {_sentence_text(s)}\n")
    return "\n".join(blocks).encode("utf-8")


def to_docx(meta: dict, data: dict) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(meta["title"], level=1)
    info = doc.add_paragraph()
    info.add_run(f"Transcrito em {meta['created'][:10]}").italic = True
    for turn in data["turns"]:
        p = doc.add_paragraph()
        p.add_run(f"{speaker_name(meta, turn['speaker'])} ({_ts(turn['start'])}): ").bold = True
        p.add_run(_turn_text(turn))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


FORMATS = {
    "txt": (to_txt, "text/plain; charset=utf-8"),
    "srt": (to_srt, "application/x-subrip; charset=utf-8"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}
